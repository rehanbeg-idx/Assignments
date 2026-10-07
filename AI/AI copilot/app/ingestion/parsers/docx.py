"""DOCX extraction using python-docx."""

from __future__ import annotations

from io import BytesIO
from pathlib import PurePosixPath

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError

from app.ingestion.exceptions import DocumentParsingError
from app.ingestion.parsers.base import DocumentParser
from app.ingestion.schemas import ParsedDocument, ParsedSection


class DOCXParser(DocumentParser):
    """Extract headings, paragraphs, and list items from DOCX files."""

    def parse(self, content: bytes, *, filename: str) -> ParsedDocument:
        try:
            document = DocxDocument(BytesIO(content))
        except PackageNotFoundError as exc:
            raise DocumentParsingError(f"Malformed or unreadable DOCX: {filename}") from exc
        except Exception as exc:  # noqa: BLE001
            raise DocumentParsingError(f"Failed to open DOCX: {filename}") from exc

        sections: list[ParsedSection] = []
        current_heading: str | None = None
        parent_heading: str | None = None
        heading_stack: list[str] = []
        section_index = 0

        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()
            if not text:
                continue

            style_name = (paragraph.style.name if paragraph.style else "") or ""
            style_lower = style_name.lower()

            if style_lower.startswith("heading"):
                level = _heading_level(style_name)
                # Maintain a simple heading stack for parent context.
                while len(heading_stack) >= level:
                    heading_stack.pop()
                parent_heading = heading_stack[-1] if heading_stack else None
                heading_stack.append(text)
                current_heading = text

                sections.append(
                    ParsedSection(
                        text=text,
                        section_index=section_index,
                        heading=current_heading,
                        parent_heading=parent_heading,
                        kind="heading",
                        metadata={"style": style_name, "heading_level": level},
                    )
                )
                section_index += 1
                continue

            kind = "list_item" if style_lower.startswith("list") else "paragraph"
            sections.append(
                ParsedSection(
                    text=text,
                    section_index=section_index,
                    heading=current_heading,
                    parent_heading=parent_heading,
                    kind=kind,
                    metadata={"style": style_name} if style_name else {},
                )
            )
            section_index += 1

        if not sections:
            raise DocumentParsingError("DOCX contains no extractable paragraphs.")

        title = (
            next((s.text for s in sections if s.kind == "heading"), None)
            or PurePosixPath(filename).stem
            or filename
        )
        raw_text = "\n\n".join(section.text for section in sections)

        return ParsedDocument(
            title=title,
            source_filename=filename,
            document_type="docx",
            sections=sections,
            raw_text=raw_text,
            metadata={"paragraph_count": len(sections)},
            warnings=[],
        )


def _heading_level(style_name: str) -> int:
    """Parse heading level from a Word style name such as 'Heading 2'."""
    parts = style_name.split()
    for part in reversed(parts):
        if part.isdigit():
            return max(1, int(part))
    return 1
